import os
import uuid
import polars as pl
from openpyxl import load_workbook


def _profile_excel_structure(path: str):
    sheets_meta = []
    wb = load_workbook(path, data_only=True, read_only=False)
    try:
        wb_f = load_workbook(path, data_only=False, read_only=False)
    except Exception:
        wb_f = None
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        hidden_cols = [
            k for k, v in ws.column_dimensions.items()
            if getattr(v, "hidden", False)
        ]
        hidden_rows = [
            k for k, v in ws.row_dimensions.items()
            if getattr(v, "hidden", False)
        ]
        merged = [str(r) for r in ws.merged_cells.ranges]
        formula_cells = []
        if wb_f is not None and sheet_name in wb_f.sheetnames:
            ws_f = wb_f[sheet_name]
            max_r = min(50, ws_f.max_row or 0)
            for row in ws_f.iter_rows(min_row=1, max_row=max_r):
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        formula_cells.append({"coord": c.coordinate, "formula": c.value})
                    if len(formula_cells) >= 50:
                        break
                if len(formula_cells) >= 50:
                    break
        sheets_meta.append({
            "sheet_name": sheet_name,
            "row_count": ws.max_row or 0,
            "column_count": ws.max_column or 0,
            "hidden_columns": hidden_cols,
            "hidden_rows": hidden_rows,
            "merged_ranges": merged,
            "formula_cells": formula_cells,
        })
    wb.close()
    if wb_f is not None:
        wb_f.close()
    return sheets_meta


def _read_csv_robust(path: str) -> pl.DataFrame:
    try:
        return pl.read_csv(path, infer_schema_length=2000, ignore_errors=True)
    except Exception:
        try:
            return pl.read_csv(path, infer_schema=False, ignore_errors=True)
        except Exception:
            return pl.read_csv(path, has_header=True, infer_schema=False,
                               truncate_ragged_lines=True)


def _read_excel_robust(path: str):
    try:
        out = pl.read_excel(path, sheet_id=0, engine="calamine")
        if isinstance(out, dict):
            return out
        return {"Sheet1": out}
    except Exception:
        pass
    try:
        df = pl.read_excel(path, engine="calamine")
        return {"Sheet1": df}
    except Exception:
        pass
    dfs = {}
    wb = load_workbook(path, data_only=True, read_only=True)
    for sn in wb.sheetnames:
        ws = wb[sn]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [
            (str(h) if h is not None else f"col_{i}")
            for i, h in enumerate(rows[0])
        ]
        seen = {}
        uniq = []
        for h in headers:
            if h in seen:
                seen[h] += 1
                uniq.append(f"{h}_{seen[h]}")
            else:
                seen[h] = 0
                uniq.append(h)
        data = [list(r) for r in rows[1:]]
        try:
            df = pl.DataFrame(data, schema=uniq, orient="row")
        except Exception:
            data_s = [[str(x) if x is not None else None for x in r] for r in data]
            df = pl.DataFrame(data_s, schema=uniq, orient="row")
        dfs[sn] = df
    wb.close()
    return dfs


def _summarize_anomalies(files):
    return {
        "merged_cells_count": sum(len(s["merged_ranges"]) for f in files for s in f["sheets"]),
        "hidden_columns_count": sum(len(s["hidden_columns"]) for f in files for s in f["sheets"]),
        "hidden_rows_count": sum(len(s["hidden_rows"]) for f in files for s in f["sheets"]),
        "formula_cells_count": sum(len(s["formula_cells"]) for f in files for s in f["sheets"]),
    }


def ingest(paths):
    session_id = str(uuid.uuid4())
    files_meta = []
    all_dfs = {}

    for idx, path in enumerate(paths):
        ext = os.path.splitext(path)[1].lower()
        fname = os.path.basename(path)
        size = os.path.getsize(path) if os.path.exists(path) else 0
        sheets_meta = []
        if ext in (".xlsx", ".xls"):
            try:
                sheets_meta = _profile_excel_structure(path)
            except Exception:
                sheets_meta = []
            dfs = _read_excel_robust(path)
        elif ext == ".csv":
            df = _read_csv_robust(path)
            dfs = {"data": df}
            sheets_meta = [{
                "sheet_name": "data",
                "row_count": df.height,
                "column_count": df.width,
                "hidden_columns": [], "hidden_rows": [],
                "merged_ranges": [], "formula_cells": [],
            }]
        else:
            raise ValueError(f"Unsupported file extension: {ext}")

        for sheet_name, df in dfs.items():
            key = f"f{idx}__{sheet_name}"
            all_dfs[key] = df
            for sm in sheets_meta:
                if sm["sheet_name"] == sheet_name:
                    sm["row_count"] = df.height
                    sm["column_count"] = df.width

        files_meta.append({
            "file_id": f"f{idx}",
            "filename": fname,
            "path": path,
            "extension": ext,
            "filesize": size,
            "sheets": sheets_meta,
        })

    manifest = {
        "session_id": session_id,
        "files": files_meta,
        "anomalies": _summarize_anomalies(files_meta),
    }
    return manifest, all_dfs